"""Extract plain text from PDF, DOCX, or plain-text files."""

from __future__ import annotations

from pathlib import Path


def extract_text(path: str | Path) -> str:
    """Return the full text content of a policy document.

    Supports:
    - .pdf  via pypdf
    - .docx via python-docx (optional dep)
    - .txt / .md and anything else — read as UTF-8 text
    """
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ImportError("pypdf is required for PDF extraction: pip install pypdf") from e

    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:
        raise ImportError(
            "python-docx is required for DOCX extraction: "
            "pip install 'guardrails-from-policy[docx]'"
        ) from e

    doc = docx.Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also pull text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n\n".join(paragraphs)
